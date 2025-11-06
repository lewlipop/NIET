#Codes for creating a new excel sheet "Vulnerabilities Summary Count" to summarise the 
# number of vulnerabiltiies recorded via severity. (PASS workbook object to function)

from openpyxl.styles import PatternFill
from collections import Counter
from excel_standard_styles import HEADER_FONT, CENTER_ALIGN, ALL_BORDERS_STYLES #Import the header font, borders and center alignment styles from excel_standard_styles.py

# Define a different header fill color for the summary sheet
SUMMARY_HEADER_FILL = PatternFill(fill_type="solid", start_color="FFD966", end_color="FFD966")  # Light orange

def add_summary_to_excel(wb):

    if "BurpSuite Scan Results" not in wb.sheetnames:
        return  # or raise an exception

    ws = wb["BurpSuite Scan Results"]
    severity_counts = Counter()

    #Count the number of vulnerabilities with the respective severity 
    for row in ws.iter_rows(min_row=2, values_only=True):
        severity = row[3]
        if severity:
            severity_counts[severity] += 1
    
    if "Vulnerabilities Summary Count" in wb.sheetnames:
        del wb["Vulnerabilities Summary Count"]
    summary_ws = wb.create_sheet("Vulnerabilities Summary Count") #Create "Vulnerabilities Summary Count" sheet

    # Define header and apply the styles
    summary_ws.append(["Severity", "Count"])
    for cell in summary_ws[1]:
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = ALL_BORDERS_STYLES
        cell.fill = SUMMARY_HEADER_FILL  # Use custom fill

    # Add the count for each severity and apply the styles 
    for i, severity in enumerate(["High", "Medium", "Low", "Information"], start=2):
        summary_ws.append([severity, severity_counts.get(severity, 0)])
        for cell in summary_ws[i]:
            cell.alignment = CENTER_ALIGN
            cell.border = ALL_BORDERS_STYLES